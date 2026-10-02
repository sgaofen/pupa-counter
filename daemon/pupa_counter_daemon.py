"""Long-lived Python worker for the desktop app.

Loads the CNN + classifier ONCE at startup, then reads one JSON command
per line on stdin and writes one JSON response per line on stdout.

Protocol (one JSON object per line, on stdin / stdout):

    request : {"id": <int>, "cmd": "detect", "imagePath": "<path>",
               "dpi": <optional int, actual scan DPI if known>}
    response: {"id": <int>, "ok": true,  "result": {<detection>}}

    request : {"id": <int>, "cmd": "export_xlsx", "path": "<out.xlsx>",
               "sheets": [{"name": str, "header": [...], "rows": [[...], ...]}]}
    response: {"id": <int>, "ok": true,  "result": {"path": "<out.xlsx>"}}

    request : {"id": <int>, "cmd": "ping"}      -> "pong"
    request : {"id": <int>, "cmd": "quit"}      -> exit(0)

Startup sentinel:

    {"ready": true, "model": "<file>", "classifier": "<file>", "trainDpi": 150, ...}

Model selection comes from model/manifest.json (file, classifier,
trainDpi, inference parameters). Env vars override individual fields:
PUPA_MODEL_MANIFEST, PUPA_MODEL_PATH, PUPA_CLF_PATH, PUPA_TRAIN_DPI,
PUPA_PEAK_THR, PUPA_MIN_DIST, PUPA_BBOX_CROP, PUPA_BBOX_HEAT_THR,
PUPA_BBOX_PAD, PUPA_CLF_PROB_THR.

DPI guard: the CNN only sees images at its training DPI. If the scan's
actual DPI differs, the image is resized to trainDpi for inference and the
peak coordinates are mapped back to the original image, so every number
returned (x, y, rank, band, sheet position) refers to the original pixels.
"""

from __future__ import annotations

import json
import os
import pickle
import struct
import sys
import traceback
from pathlib import Path

import cv2
import numpy as np
import torch

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from pupa_counter import (  # type: ignore
    TinyUNet,
    predict_heatmap,
    extract_peaks,
    _peak_features,
    pick_device,
    device_description,
)

import stage2  # type: ignore

try:
    import sheet_detect
except Exception:  # pragma: no cover - sheet detection is optional
    sheet_detect = None


STANDARD_DPIS = (75, 100, 150, 200, 240, 300, 400, 600, 1200)
# Scan-bed widths / heights we know about, in inches. The WIA script
# requests an A4 window (8.27 x 11.69); ImageCaptureCore returns the full
# LiDE 300 bed (8.5 x 11.69).
BED_WIDTHS_IN = (8.27, 8.5)
BED_HEIGHTS_IN = (11.69, 11.7)

SHEET_KEYS = ("found", "corners", "angleDeg", "lengthPx", "widthPx", "truncatedTop",
              "truncatedBottom", "confidence", "method")

SUSPECT_HEAT_THR = 0.30
SUSPECT_PROB_MIN = 0.20
SUSPECT_MAX = 60


def _resolve(env_name: str, default: Path) -> Path:
    raw = os.environ.get(env_name)
    if raw:
        return Path(raw).expanduser().resolve()
    return default


def load_manifest() -> dict:
    path = _resolve("PUPA_MODEL_MANIFEST", HERE / "model" / "manifest.json")
    data: dict = {}
    if path.exists():
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
    data.setdefault("file", "pupa_counter_lide300.pt")
    data.setdefault("classifier", "peak_filter_clf_lide300.pkl")
    data.setdefault("trainDpi", 150)
    data.setdefault("inference", {})
    data["_path"] = str(path)
    return data


MANIFEST = load_manifest()
_INF = MANIFEST.get("inference", {})


def _envf(name: str, default: float) -> float:
    raw = os.environ.get(name)
    return float(raw) if raw not in (None, "") else float(default)


_PEAK_THR = _envf("PUPA_PEAK_THR", _INF.get("peakThr", 0.50))
_MIN_DIST = int(_envf("PUPA_MIN_DIST", _INF.get("minDist", 3)))
_BBOX_CROP = os.environ.get("PUPA_BBOX_CROP", "1" if _INF.get("bboxCrop", True) else "0") == "1"
_BBOX_HEAT_THR = _envf("PUPA_BBOX_HEAT_THR", _INF.get("bboxHeatThr", 0.40))
_BBOX_PAD = int(_envf("PUPA_BBOX_PAD", _INF.get("bboxPad", 50)))
_CLF_PROB_THR = _envf("PUPA_CLF_PROB_THR", _INF.get("clfProbThr", 0.50))
_STAGE2 = None  # loaded in main() when the manifest has a "stage2" block
_TILING = _INF.get("tiling", "padded")
_EXCLUDE_BORDER = bool(_INF.get("excludeBorder", True))
_SKIP_BLANK = bool(_INF.get("skipBlankTiles", False))
TRAIN_DPI = int(_envf("PUPA_TRAIN_DPI", MANIFEST.get("trainDpi", 150)))


def _heat_bbox(heatmap: np.ndarray, blob_thr: float, pad: int):
    """Bounding box of the high-confidence blob (crops blank scanner bed)."""
    above = heatmap > blob_thr
    rows = np.any(above, axis=1)
    if not rows.any():
        return None
    cols = np.any(above, axis=0)
    H, W = heatmap.shape
    y0 = int(rows.argmax())
    y1_inv = int(rows[::-1].argmax())
    x0 = int(cols.argmax())
    x1_inv = int(cols[::-1].argmax())
    return (
        max(0, x0 - pad),
        max(0, y0 - pad),
        min(W, W - x1_inv + pad),
        min(H, H - y1_inv + pad),
    )


# --- DPI resolution -----------------------------------------------------------

def _snap_dpi(v: float, tol: float = 0.04):
    best = min(STANDARD_DPIS, key=lambda d: abs(d - v))
    return best if abs(best - v) <= tol * best else None


def infer_dpi_from_size(w: int, h: int):
    """Guess DPI from pixel size, assuming a full-bed / A4 scan window."""
    for bw in BED_WIDTHS_IN:
        d = _snap_dpi(w / bw)
        if d is None:
            continue
        if any(abs(h - bh * d) <= 0.03 * bh * d for bh in BED_HEIGHTS_IN):
            return d
    return None


def png_phys_dpi(path: Path):
    """Read the pHYs chunk of a PNG (pixels per metre) -> DPI, or None."""
    try:
        with open(path, "rb") as f:
            if f.read(8) != b"\x89PNG\r\n\x1a\n":
                return None
            while True:
                head = f.read(8)
                if len(head) < 8:
                    return None
                length, ctype = struct.unpack(">I4s", head)
                if ctype == b"pHYs":
                    data = f.read(9)
                    ppx, ppy, unit = struct.unpack(">IIB", data)
                    if unit != 1 or ppx == 0:
                        return None
                    return round(ppx * 0.0254)
                if ctype == b"IDAT":
                    return None
                f.seek(length + 4, 1)
    except Exception:
        return None


def resolve_dpi(path: Path, w: int, h: int, requested):
    if requested:
        try:
            v = int(round(float(requested)))
            if v > 0:
                return v, "scanner"
        except (TypeError, ValueError):
            pass
    d = infer_dpi_from_size(w, h)
    if d:
        return d, "pixel-size"
    d = png_phys_dpi(path)
    if d and 50 <= d <= 4800:
        return d, "png-metadata"
    return TRAIN_DPI, "assumed"


# --- Rank / band (same arithmetic as pupa_counter.render_annotated) ----------

def compute_ranks(peaks):
    counts = {"top_5_pct": 0, "rank_5_to_25_pct": 0, "middle_50_pct": 0,
              "bottom_25_pct": 0, "y_min_of_pupae": None, "y_max_of_pupae": None}
    per = []
    if not peaks:
        return counts, per
    ys = np.array([y for _, y in peaks])
    y_min, y_max = int(ys.min()), int(ys.max())
    y_range = max(1, y_max - y_min)
    counts["y_min_of_pupae"], counts["y_max_of_pupae"] = y_min, y_max

    def line_y(p: float) -> int:
        return int(y_max - (p / 100.0) * y_range)

    y5, y25, y75 = line_y(5), line_y(25), line_y(75)
    for x, y in peaks:
        rank = (y_max - y) / y_range * 100.0
        if y >= y5:
            band = "0-5%"; counts["top_5_pct"] += 1
        elif y >= y25:
            band = "5-25%"; counts["rank_5_to_25_pct"] += 1
        elif y >= y75:
            band = "25-75%"; counts["middle_50_pct"] += 1
        else:
            band = "75-100%"; counts["bottom_25_pct"] += 1
        per.append((round(rank, 2), band))
    return counts, per


# --- Worker -------------------------------------------------------------------

def main() -> None:
    try:
        model_path = _resolve("PUPA_MODEL_PATH", HERE / "model" / MANIFEST["file"])
        clf_name = MANIFEST.get("classifier")
        clf_path = (_resolve("PUPA_CLF_PATH", HERE / "model" / (clf_name or ""))
                    if clf_name or os.environ.get("PUPA_CLF_PATH") else None)
        if not model_path.exists():
            raise FileNotFoundError(f"model weights not found: {model_path}")
        device = pick_device()
        model = TinyUNet().to(device)
        model.load_state_dict(torch.load(model_path, map_location=device))
        model.eval()
        classifier = None
        if clf_path is not None and clf_path.exists():
            with open(clf_path, "rb") as f:
                classifier = pickle.load(f)
        global _STAGE2
        st2 = MANIFEST.get("stage2")
        if st2:
            with open(_resolve("PUPA_STAGE2_PATH", HERE / "model" / st2["file"]), "rb") as f:
                bundle = pickle.load(f)
            for k in ("cand_thr", "t_keep", "t2", "dist"):
                if k in st2:
                    bundle[k] = st2[k]
            _STAGE2 = bundle
    except Exception as exc:
        sys.stdout.write(json.dumps({
            "ready": False, "stage": "startup",
            "error": f"{type(exc).__name__}: {exc}",
            "traceback": traceback.format_exc(),
        }) + "\n")
        sys.stdout.flush()
        sys.exit(2)

    sys.stdout.write(json.dumps({
        "ready": True,
        "device": str(device),
        "deviceName": device_description(device),
        "model": model_path.name,
        "modelName": MANIFEST.get("name", model_path.stem),
        "classifier": (MANIFEST["stage2"]["file"] if _STAGE2 is not None
                       else clf_path.name if classifier is not None else None),
        "trainDpi": TRAIN_DPI,
        "manifest": MANIFEST.get("_path"),
        "sheetDetector": getattr(sheet_detect, "__name__", None) and (
            "stub" if "STUB" in (sheet_detect.__doc__ or "").upper() else "real"),
        "config": {
            "peak_thr": _PEAK_THR, "min_dist": _MIN_DIST,
            "bbox_crop": _BBOX_CROP, "clf_prob_thr": _CLF_PROB_THR,
        },
    }) + "\n")
    sys.stdout.flush()

    for raw in sys.stdin:
        line = raw.strip()
        if not line:
            continue
        try:
            req = json.loads(line)
        except json.JSONDecodeError as exc:
            _respond({"id": None, "ok": False, "error": f"bad JSON: {exc}"})
            continue
        rid = req.get("id")
        cmd = req.get("cmd")
        try:
            if cmd == "ping":
                _respond({"id": rid, "ok": True, "result": "pong"})
            elif cmd == "detect":
                if req.get("analysisPath"):
                    result = detect_pair(req["imagePath"], req["analysisPath"], model, device,
                                         classifier, model_path.name, req.get("dpi"),
                                         req.get("analysisDpi"))
                else:
                    result = detect(req["imagePath"], model, device, classifier,
                                    model_path.name, req.get("dpi"))
                _respond({"id": rid, "ok": True, "result": result})
            elif cmd == "export_xlsx":
                _respond({"id": rid, "ok": True, "result": export_xlsx(req)})
            elif cmd == "quit":
                _respond({"id": rid, "ok": True})
                return
            else:
                _respond({"id": rid, "ok": False, "error": f"unknown cmd: {cmd!r}"})
        except Exception as exc:
            _respond({
                "id": rid, "ok": False,
                "error": f"{type(exc).__name__}: {exc}",
                "traceback": traceback.format_exc(),
            })


def _respond(obj: dict) -> None:
    sys.stdout.write(json.dumps(obj, separators=(",", ":")) + "\n")
    sys.stdout.flush()


def run_model(img_rgb: np.ndarray, model, device, classifier):
    """Run the ship pipeline on an image already at TRAIN_DPI.

    Returns (kept, suspects) in the coordinates of `img_rgb`:
      kept     = [(x, y, heat, prob|None)]
      suspects = [(x, y, heat, prob|None, reason)]
    """
    heatmap = predict_heatmap(model, img_rgb, device, tiling=_TILING, skip_blank=_SKIP_BLANK)
    if _STAGE2 is not None:
        pts, sus = stage2.run(heatmap, img_rgb, _STAGE2)
        H, W = heatmap.shape
        at = lambda x, y: float(heatmap[min(H - 1, max(0, int(round(y)))), min(W - 1, max(0, int(round(x))))])
        kept = [(x, y, at(x, y), p) for x, y, p in pts]
        suspects = [(x, y, at(x, y), p, "low-confidence") for x, y, p in sus]
        suspects.sort(key=lambda t: -t[3])
        return kept, suspects[:SUSPECT_MAX]
    raw_peaks = extract_peaks(heatmap, threshold=_PEAK_THR, min_dist=_MIN_DIST,
                              exclude_border=_EXCLUDE_BORDER)
    bbox = _heat_bbox(heatmap, _BBOX_HEAT_THR, _BBOX_PAD) if _BBOX_CROP else None
    if bbox is not None:
        x0, y0, x1, y1 = bbox
        raw_peaks = [(x, y) for x, y in raw_peaks if x0 <= x < x1 and y0 <= y < y1]

    if classifier is not None and raw_peaks:
        # Same features / threshold as pupa_counter.filter_false_positives,
        # computed once so the probabilities can be reported too.
        probs = classifier.predict_proba(_peak_features(raw_peaks, heatmap, img_rgb))[:, 1]
    else:
        probs = [None] * len(raw_peaks)

    kept, suspects = [], []
    for (x, y), pr in zip(raw_peaks, probs):
        heat = float(heatmap[y, x])
        if pr is None or pr >= _CLF_PROB_THR:
            kept.append((x, y, heat, None if pr is None else float(pr)))
        elif pr >= SUSPECT_PROB_MIN:
            suspects.append((x, y, heat, float(pr), "classifier-rejected"))

    # Weak heatmap peaks the threshold dropped: possible misses (often the
    # dark / out-of-focus pupae). Only inside the pupa region.
    raw_set = np.array([(x, y) for x, y in raw_peaks], dtype=np.float32).reshape(-1, 2)
    weak = extract_peaks(heatmap, threshold=SUSPECT_HEAT_THR, min_dist=_MIN_DIST,
                         exclude_border=_EXCLUDE_BORDER)
    far = max(4, 2 * _MIN_DIST)
    for x, y in weak:
        heat = float(heatmap[y, x])
        if heat >= _PEAK_THR:
            continue
        if bbox is not None:
            x0, y0, x1, y1 = bbox
            if not (x0 <= x < x1 and y0 <= y < y1):
                continue
        if len(raw_set) and np.min(np.hypot(raw_set[:, 0] - x, raw_set[:, 1] - y)) < far:
            continue
        suspects.append((x, y, heat, None, "weak-signal"))
    suspects.sort(key=lambda s: -(s[3] if s[3] is not None else s[2]))
    return kept, suspects[:SUSPECT_MAX]


def detect(image_path: str, model, device, classifier, model_name: str, dpi=None) -> dict:
    path = Path(image_path)
    img_bgr = cv2.imread(str(path), cv2.IMREAD_COLOR)
    if img_bgr is None:
        raise FileNotFoundError(f"could not read image: {image_path}")
    h, w = img_bgr.shape[:2]
    image_dpi, dpi_source = resolve_dpi(path, w, h, dpi)

    scale = TRAIN_DPI / float(image_dpi)
    if abs(scale - 1.0) > 0.02:
        mw, mh = max(1, int(round(w * scale))), max(1, int(round(h * scale)))
        interp = cv2.INTER_AREA if scale < 1 else cv2.INTER_CUBIC
        model_bgr = cv2.resize(img_bgr, (mw, mh), interpolation=interp)
        sx, sy = w / mw, h / mh
    else:
        scale = 1.0
        model_bgr = img_bgr
        sx = sy = 1.0

    kept, suspects = run_model(cv2.cvtColor(model_bgr, cv2.COLOR_BGR2RGB),
                               model, device, classifier)

    def to_orig(x, y):
        if sx == 1.0 and sy == 1.0:
            return int(x), int(y)
        ox = int(round((x + 0.5) * sx - 0.5))
        oy = int(round((y + 0.5) * sy - 0.5))
        return min(max(ox, 0), w - 1), min(max(oy, 0), h - 1)

    peaks = [to_orig(x, y) for x, y, _, _ in kept]
    counts, ranks = compute_ranks(peaks)

    sheet = {"found": False, "corners": [], "confidence": 0.0, "method": "unavailable"}
    sheet_pct = [None] * len(peaks)
    if sheet_detect is not None:
        try:
            # The sheet detector's size priors are in TRAIN_DPI (150 DPI)
            # pixels, so run it on the same resized image the model saw and
            # map the corners back to original pixels.
            raw = sheet_detect.detect_sheet(model_bgr, pupae_xy=[(x, y) for x, y, _, _ in kept])
            sheet = {k: raw.get(k) for k in SHEET_KEYS if k in raw}
            sheet = json.loads(json.dumps(sheet, default=float))
            if sheet.get("corners"):
                sheet["corners"] = [[round((cx + 0.5) * sx - 0.5, 1), round((cy + 0.5) * sy - 0.5, 1)]
                                    for cx, cy in sheet["corners"]]
            for k in ("lengthPx", "widthPx"):
                if isinstance(sheet.get(k), (int, float)):
                    sheet[k] = round(sheet[k] * sx, 1)
            if sheet.get("found") and len(sheet.get("corners") or []) == 4:
                sheet_pct = [float(sheet_detect.sheet_relative_pct(sheet["corners"], x, y))
                             for x, y in peaks]
        except Exception as exc:
            sheet = {"found": False, "corners": [], "confidence": 0.0,
                     "method": "error", "error": f"{type(exc).__name__}: {exc}"}

    pupae = []
    for i, ((x, y), (rank, band), (_, _, heat, prob)) in enumerate(zip(peaks, ranks, kept)):
        pupae.append({
            "index": i + 1, "x": x, "y": y,
            "rankPct": rank, "band": band,
            "sheetPct": None if sheet_pct[i] is None else round(sheet_pct[i], 2),
            "score": round(heat, 3),
            "prob": None if prob is None else round(prob, 3),
            "source": "cnn",
        })
    sus = []
    for x, y, heat, prob, reason in suspects:
        ox, oy = to_orig(x, y)
        sus.append({"x": ox, "y": oy, "score": round(heat, 3),
                    "prob": None if prob is None else round(prob, 3), "reason": reason})

    return {
        "imagePath": str(path),
        "imageWidth": w,
        "imageHeight": h,
        "imageDpi": image_dpi,
        "imageDpiSource": dpi_source,
        "trainDpi": TRAIN_DPI,
        "inferenceScale": round(scale, 4),
        "modelVersion": model_name,
        "pupae": pupae,
        "suspects": sus,
        "sheet": sheet,
        "counts": {
            "total": len(peaks),
            "top5Pct": counts["top_5_pct"],
            "rank5To25": counts["rank_5_to_25_pct"],
            "middle50": counts["middle_50_pct"],
            "bottom25": counts["bottom_25_pct"],
        },
        "yMin": counts["y_min_of_pupae"],
        "yMax": counts["y_max_of_pupae"],
    }


def align_shift(analysis_bgr: np.ndarray, display_bgr: np.ndarray):
    """Offset of the display scan relative to the analysis scan, measured on
    the analysis pixel grid (phase correlation of the display image resized
    to the analysis size). Two passes on a flatbed only differ by the
    carriage start position, i.e. a small translation."""
    ha, wa = analysis_bgr.shape[:2]
    a = cv2.cvtColor(analysis_bgr, cv2.COLOR_BGR2GRAY).astype(np.float32)
    d = cv2.cvtColor(cv2.resize(display_bgr, (wa, ha), interpolation=cv2.INTER_AREA),
                     cv2.COLOR_BGR2GRAY).astype(np.float32)
    win = cv2.createHanningWindow((wa, ha), cv2.CV_32F)
    (dx, dy), resp = cv2.phaseCorrelate(a, d, win)
    return float(dx), float(dy), float(resp)


def detect_pair(display_path: str, analysis_path: str, model, device, classifier,
                model_name: str, dpi=None, analysis_dpi=None) -> dict:
    """Detect on the native 150-DPI analysis scan, then copy the detections
    onto the high-DPI display scan (aligned + scaled)."""
    res = detect(analysis_path, model, device, classifier, model_name,
                 analysis_dpi or TRAIN_DPI)
    disp = cv2.imread(str(display_path), cv2.IMREAD_COLOR)
    if disp is None:
        raise FileNotFoundError(f"could not read image: {display_path}")
    ana = cv2.imread(str(analysis_path), cv2.IMREAD_COLOR)
    hd, wd = disp.shape[:2]
    ha, wa = ana.shape[:2]
    fx, fy = wd / wa, hd / ha
    dx, dy, resp = align_shift(ana, disp)
    warn = None
    if resp < 0.05 or abs(dx) > 40 or abs(dy) > 40:
        warn = f"alignment unreliable (shift {dx:.1f},{dy:.1f} px, response {resp:.3f}); assuming no shift"
        dx = dy = 0.0

    def m(x, y):
        return (min(max((x + dx + 0.5) * fx - 0.5, 0), wd - 1),
                min(max((y + dy + 0.5) * fy - 0.5, 0), hd - 1))

    for p in res["pupae"]:
        p["x"], p["y"] = (round(v, 1) for v in m(p["x"], p["y"]))
    for p in res["suspects"]:
        p["x"], p["y"] = (round(v, 1) for v in m(p["x"], p["y"]))
    sh = res.get("sheet") or {}
    if sh.get("corners"):
        sh["corners"] = [[round(c, 1) for c in m(cx, cy)] for cx, cy in sh["corners"]]
    for k in ("lengthPx", "widthPx"):
        if isinstance(sh.get(k), (int, float)):
            sh[k] = round(sh[k] * fx, 1)
    if res.get("yMin") is not None:
        res["yMin"] = round(m(0, res["yMin"])[1], 1)
        res["yMax"] = round(m(0, res["yMax"])[1], 1)
    disp_dpi, disp_src = resolve_dpi(Path(display_path), wd, hd, dpi)
    res.update({
        "imagePath": str(display_path), "imageWidth": wd, "imageHeight": hd,
        "imageDpi": disp_dpi, "imageDpiSource": disp_src,
        "analysis": {"path": str(analysis_path), "width": wa, "height": ha,
                     "dpi": res.get("imageDpi"), "scale": [round(fx, 4), round(fy, 4)],
                     "shiftPx": [round(dx, 2), round(dy, 2)], "alignResponse": round(resp, 3),
                     "warning": warn},
    })
    return res


def export_xlsx(req: dict) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    out = Path(req["path"])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for sh in req.get("sheets", []):
        ws = wb.create_sheet(str(sh.get("name", "Sheet"))[:31])
        header = sh.get("header", [])
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in sh.get("rows", []):
            ws.append(row)
        ws.freeze_panes = "A2"
        for i, name in enumerate(header, start=1):
            col = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[col].width = max(8, min(40, len(str(name)) + 2))
    wb.save(out)
    return {"path": str(out)}


if __name__ == "__main__":
    main()
